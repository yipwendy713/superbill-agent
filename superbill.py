"""Superbill logic: turn a PT session summary into coded line items.

    python superbill.py
"""
import html
import json
from datetime import date, datetime
from pathlib import Path

# Superbills are saved in a "superbills" folder next to this file.
OUTPUT_DIR = Path(__file__).resolve().parent / "superbills"

# Placeholder practice info ).
PRACTICE = {

    "name": "Athletic PT, Inc",
    "address": "1234 Main St Los Angeles, 90014",
    "tax_id": "12-345-3910",
    "clinician": "Lisa Taylor, PT, DPT",
    "license_number": "PT 12345",
    "npi": "123 789 3418",
}

# Fee schedule from the practice's superbill. None = no price on file.
FEE_SCHEDULE = {
    "97161": {"description": "PT Eval Low", "timed": False, "unit_price": 85.00},
    "97162": {"description": "PT Eval Mod", "timed": False, "unit_price": 85.00},
    "97163": {"description": "PT Eval High", "timed": False, "unit_price": 85.00},
    "97140": {"description": "Manual Therapy", "timed": True, "unit_price": 50.00},
    "97530": {"description": "Kinetic Therapy", "timed": True, "unit_price": None},
    "97110": {"description": "Therapeutic Exercise", "timed": True, "unit_price": 50.00},
    "97112": {"description": "Neuro Re Ed", "timed": True, "unit_price": None},
    "97035": {"description": "Ultrasound", "timed": True, "unit_price": 25.00},
    "97032": {"description": "E Stim Attended", "timed": True, "unit_price": 25.00},
    "97014": {"description": "E Stim Unattended", "timed": False, "unit_price": None},
}

EVAL_CODES = {"low": "97161", "moderate": "97162", "high": "97163"}

# 77 common PT diagnoses: label -> (ICD-10 code, description). See diagnoses.py.
from diagnoses import DIAGNOSES  # noqa: E402

MODIFIER = "GP"  # physical therapy services


def units_from_minutes(minutes: int) -> int:
    """15-minute units: each full 15 minutes is one unit, and a leftover
    of 8+ minutes adds one more. Simplified per-service rule; Medicare's
    total-time method can give different results across multiple codes."""
    if minutes < 8:
        return 0
    full, remainder = divmod(minutes, 15)
    return full + (1 if remainder >= 8 else 0)


def build_superbill(patient_name, service_date, diagnosis, eval_level,
                    timed_minutes, estim_unattended):
    """timed_minutes: dict of CPT code -> minutes, e.g. {"97140": 30}"""
    lines, warnings = [], []

    def add_line(code, units, minutes=None):
        fee = FEE_SCHEDULE[code]
        price = fee["unit_price"]
        if price is None:
            warnings.append(f"No price on file for {code} {fee['description']}; needs review.")
        lines.append({
            "cpt_code": code,
            "description": fee["description"],
            "modifier": MODIFIER,
            "minutes": minutes,
            "units": units,
            "unit_price": price,
            "cost": round(price * units, 2) if price is not None else None,
        })

    if eval_level in EVAL_CODES:
        add_line(EVAL_CODES[eval_level], 1)

    for code, minutes in timed_minutes.items():
        minutes = minutes or 0
        if minutes == 0:
            continue
        units = units_from_minutes(minutes)
        if units == 0:
            warnings.append(f"{minutes} min of {FEE_SCHEDULE[code]['description']} "
                            "is under 8 minutes, so it is not billed.")
            continue
        add_line(code, units, minutes)

    if estim_unattended:
        add_line("97014", 1)

    icd10, dx_description = DIAGNOSES[diagnosis]
    total = round(sum(line["cost"] or 0 for line in lines), 2)
    return {
        "document_type": "superbill",
        "practice": PRACTICE,
        "patient": {"name": patient_name},
        "encounter": {"date_of_service": service_date.isoformat(), "place_of_service": "11"},
        "diagnoses": [{"icd10_code": icd10, "description": dx_description}],
        "line_items": lines,
        "totals": {"total_charges": total, "amount_paid": 0.00, "balance_due": total},
        "status": "draft",  # the PT approves before it goes to the patient
        "warnings": warnings,
    }


def spoken_summary(bill) -> str:
    """Short summary for the agent to read back. Leaves out the patient's name."""
    parts = [f"{l['units']} unit{'s' if l['units'] != 1 else ''} of {l['description']}"
             for l in bill["line_items"]]
    text = "Draft superbill: " + ", ".join(parts) if parts else "No billable services."
    text += f". Total ${bill['totals']['total_charges']:.2f}."
    if bill["warnings"]:
        text += " Needs review: " + " ".join(bill["warnings"])
    return text


def save_superbill(bill, folder=OUTPUT_DIR) -> Path:
    """Save as a Word document (for the PT), plus JSON and HTML.
    File names use a timestamp, not the patient's name."""
    out = Path(folder)
    out.mkdir(exist_ok=True)
    stem = "superbill_" + datetime.now().strftime("%Y%m%d_%H%M%S")
    base, i = stem, 2
    while (out / f"{stem}.json").exists():  # never overwrite an earlier superbill
        stem, i = f"{base}_{i}", i + 1
    (out / f"{stem}.json").write_text(json.dumps(bill, indent=2))
    (out / f"{stem}.html").write_text(render_html(bill))
    try:
        docx_path = out / f"{stem}.docx"
        render_docx(bill, docx_path)
        return docx_path
    except ImportError:
        print("python-docx is not installed; run `uv add python-docx` to get Word output.")
        return out / f"{stem}.html"


def _money(value):
    return f"${value:,.2f}" if value is not None else "—"


def _friendly_date(iso_date):
    d = date.fromisoformat(iso_date)
    return f"{d:%B} {d.day}, {d.year}"


def render_docx(bill, path):
    """Word version laid out like a practice's superbill, so the PT can
    edit it and use Save as PDF."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    accent = RGBColor(0xE8, 0x5D, 0x4F)
    p = bill["practice"]
    dx = bill["diagnoses"][0]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    run = title.add_run(p["name"])
    run.bold, run.font.size = True, Pt(18)
    heading = doc.add_paragraph()
    run = heading.add_run("SUPERBILL")
    run.font.size, run.font.color.rgb = Pt(14), accent
    draft = doc.add_paragraph()
    run = draft.add_run("DRAFT: awaiting clinician review and approval")
    run.bold, run.font.color.rgb = True, RGBColor(0xB4, 0x53, 0x09)

    doc.add_paragraph(
        f"Treating clinician: {p['clinician']}\n"
        f"PT License #: {p['license_number']}\n"
        f"NPI #: {p['npi']}\n"
        f"Tax ID: {p['tax_id']}\n"
        f"{p['address']}"
    )

    info = doc.add_table(rows=0, cols=2)
    for label, value in [
        ("Patient", bill["patient"]["name"]),
        ("Date of service", _friendly_date(bill["encounter"]["date_of_service"])),
        ("Place of service", bill["encounter"]["place_of_service"] + " (office)"),
        ("Diagnosis", f"{dx['description']} {dx['icd10_code']}"),
    ]:
        cells = info.add_row().cells
        cells[0].text, cells[1].text = label, value
        cells[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells,
                          ["CPT Code", "Minutes", "Units", "Unit Price", "Cost"]):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for line in bill["line_items"]:
        cells = table.add_row().cells
        cells[0].text = f"{line['cpt_code']} {line['description']} {line['modifier']}"
        cells[1].text = str(line["minutes"] or "")
        cells[2].text = str(line["units"])
        cells[3].text = _money(line["unit_price"])
        cells[4].text = _money(line["cost"])
    for label, value in [("Total", bill["totals"]["total_charges"]),
                         ("Amount Paid", bill["totals"]["amount_paid"])]:
        cells = table.add_row().cells
        cells[3].text, cells[4].text = label, _money(value)
        cells[3].paragraphs[0].runs[0].bold = True
        cells[4].paragraphs[0].runs[0].bold = True
    for row in table.rows:
        for i in (1, 2, 3, 4):
            row.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    if bill["warnings"]:
        doc.add_paragraph()
        note = doc.add_paragraph()
        run = note.add_run("Needs review before approval")
        run.bold, run.font.color.rgb = True, RGBColor(0xB4, 0x53, 0x09)
        for warning in bill["warnings"]:
            doc.add_paragraph(warning, style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph("Clinician signature: ____________________________    Date: ______________")
    doc.save(path)


def render_html(bill) -> str:
    e = html.escape
    p = bill["practice"]
    rows = "".join(
        f"<tr><td>{l['cpt_code']} {e(l['description'])} {l['modifier']}</td>"
        f"<td>{l['minutes'] or ''}</td><td>{l['units']}</td>"
        f"<td>{'$%.2f' % l['unit_price'] if l['unit_price'] is not None else '—'}</td>"
        f"<td>{'$%.2f' % l['cost'] if l['cost'] is not None else '—'}</td></tr>"
        for l in bill["line_items"])
    warnings = "".join(f"<li>{e(w)}</li>" for w in bill["warnings"])
    dx = bill["diagnoses"][0]
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>Superbill (draft)</title>
<style>body{{font-family:system-ui,sans-serif;max-width:760px;margin:40px auto;color:#222}}
table{{width:100%;border-collapse:collapse}}td,th{{border-bottom:1px solid #ddd;padding:8px;text-align:left}}
th{{background:#e85d4f;color:#fff}}.draft{{color:#b45309;font-weight:600}}.warn{{background:#fff7ed;padding:8px 24px}}</style>
</head><body>
<h1>{e(p['name'])}</h1><p class="draft">DRAFT: awaiting clinician approval</p>
<p>{e(p['clinician'])} · License {e(p['license_number'])} · NPI {e(p['npi'])} · Tax ID {e(p['tax_id'])}<br>{e(p['address'])}</p>
<p><b>Patient:</b> {e(bill['patient']['name'])}<br><b>Date of service:</b> {bill['encounter']['date_of_service']}
 · <b>Place of service:</b> {bill['encounter']['place_of_service']}<br>
<b>Diagnosis:</b> {e(dx['description'])} ({dx['icd10_code']})</p>
<table><tr><th>CPT code</th><th>Minutes</th><th>Units</th><th>Unit price</th><th>Cost</th></tr>{rows}
<tr><td colspan="4"><b>Total</b></td><td><b>${bill['totals']['total_charges']:.2f}</b></td></tr></table>
{f'<div class="warn"><b>Needs review</b><ul>{warnings}</ul></div>' if warnings else ''}
</body></html>"""


if __name__ == "__main__":
    # Quick test with a fake patient.
    bill = build_superbill(
        patient_name="Jane Doe",
        service_date=date.today(),
        diagnosis="Pain in right knee (M25.561)",
        eval_level="none",
        timed_minutes={"97140": 25, "97110": 30, "97112": 6},
        estim_unattended=False,
    )
    print(json.dumps(bill, indent=2))
    print("\n" + spoken_summary(bill))
    print("\nSaved:", save_superbill(bill))
